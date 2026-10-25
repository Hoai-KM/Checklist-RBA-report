# -*- coding: utf-8 -*-
"""
Động cơ đọc/ghi file Excel báo cáo HR audit (giữ nguyên toàn bộ định dạng).

Sử dụng openpyxl: style, merged cells, column widths, formulas, data validation,
external links (ví dụ VLOOKUP trỏ [2]Table!...) được bảo toàn khi load + save.
Rich text (chữ đậm header trong cột L) được ghi đúng định dạng file mẫu đã duyệt.

Cấu trúc workbook (chuẩn Samsung SEVT HR audit):
- 'Company Profile'        : thông tin nhà máy, conformance rate, interview score...
- 'General Information'    : thông tin đợt audit (loại, ngày, team, số Worker interview...)
- 'Checklist'              : 52 dòng tiêu chí (row 6-57); cột D = rating, cột L = evidence
- 'Reference table','Table': bảng tra điểm (không thay đổi)
"""
import copy
import os
import re
import tempfile
import zipfile

import openpyxl
from openpyxl.cell.rich_text import CellRichText

try:  # chạy trực tiếp (script) hoặc như package
    from knowledge_base import ROW_ITEMS, APP_VERSION
except ImportError:  # pragma: no cover
    from .knowledge_base import ROW_ITEMS, APP_VERSION

try:
    from rewrite import (
        HEADER_EN, HEADER_VI, REFS_EN, REFS_VI,
        PLACEHOLDER, build_content, ensure_section_gap, is_placeholder,
        missing_input_items, normalize_bilingual_block, parse_bilingual,
        purify_findings, scan_language_mixing, separate_languages, to_rich_text,
        _norm_unicode,
    )
except ImportError:  # pragma: no cover
    from .rewrite import (
        HEADER_EN, HEADER_VI, REFS_EN, REFS_VI,
        PLACEHOLDER, build_content, ensure_section_gap, is_placeholder,
        missing_input_items, normalize_bilingual_block, parse_bilingual,
        purify_findings, scan_language_mixing, separate_languages, to_rich_text,
        _norm_unicode,
    )

# 4 header khối song ngữ trong cột L (in đậm theo định dạng file mẫu đã duyệt)
BILINGUAL_HEADERS = (HEADER_EN, REFS_EN, HEADER_VI, REFS_VI)

CHECKLIST = "Checklist"
GENERAL_INFO = "General Information"
COMPANY_PROFILE = "Company Profile"

# Ô trong Company Profile được "reset" khi sang kỳ audit mới (giá trị kỳ cũ)
PROFILE_RESET_CELLS = ["G3", "H4", "H5", "H6", "C6", "C7", "C8"]

ALLOWED_RATINGS = ["Conformance", "Minor", "Major", "Priority", "N/A"]


def _repair_custom_xml(path):
    """
    Một số file .xlsx tạo bởi phần mềm của đối tác/vendor (vd NSC) chứa
    docProps/custom.xml với <property> thiếu thuộc tính 'name' -> openpyxl vỡ khi
    đọc custom document properties (TypeError). Hàm này tạo bản sao tạm đã được
    sửa chữa (thêm 'name' cho mọi property) và trả về đường dẫn file tạm.
    Trả về None nếu không cần sửa hoặc không sửa được.
    """
    try:
        with zipfile.ZipFile(path, "r") as zin:
            names = zin.namelist()
            if "docProps/custom.xml" not in names:
                return None
            data = zin.read("docProps/custom.xml").decode("utf-8", errors="replace")
            counter = [0]

            def _add_name(m):
                tag = m.group(0)
                if "name=" in tag:
                    return tag
                counter[0] += 1
                return tag[:-1] + f' name="AuditProp{counter[0]}" ' + tag[-1]

            new_data = re.sub(r"<property\s+[^>]*?>", _add_name, data)
            if new_data == data:
                return None  # không có property lỗi
            fd, tmp = tempfile.mkstemp(prefix="hr_audit_", suffix=".xlsx")
            os.close(fd)
            with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
                for n in names:
                    if n == "docProps/custom.xml":
                        zout.writestr(n, new_data.encode("utf-8"))
                    else:
                        zout.writestr(n, zin.read(n))
            return tmp
    except Exception:
        return None


class Report:
    """Bao bọc workbook báo cáo audit."""

    def __init__(self, path):
        self.path = path
        self.rich_text_ok = True
        self._tmp = None
        # Kiểm tra file có đúng là .xlsx thật không (zip chứa workbook)
        if not zipfile.is_zipfile(path):
            raise ValueError(
                "File không phải Excel hợp lệ (*.xlsx).\n"
                "Nguyên nhân thường gặp: file bị đổi đuôi .xlsx từ file khác "
                "(.txt/.html/.xls cũ), tải xuống chưa hoàn tất, hoặc file bị hỏng.\n"
                "→ Hãy dùng đúng file .xlsx gốc (ví dụ file đầu vào trong thư mục dự án)."
            )
        try:
            self.wb = openpyxl.load_workbook(path, data_only=False, rich_text=True)
        except Exception as e1:
            # File khác lạ có thể chứa custom property hỏng -> thử sửa chữa
            repaired = _repair_custom_xml(path)
            if repaired is None:
                raise ValueError(
                    "Không đọc được nội dung file Excel.\n"
                    "File có thể bị hỏng hoặc do phần mềm khác tạo ra mà app chưa "
                    "hỗ trợ.\nLiên hệ người quản trị app kèm file để kiểm tra."
                ) from e1
            self._tmp = repaired
            try:
                try:
                    self.wb = openpyxl.load_workbook(repaired, data_only=False,
                                                     rich_text=True)
                except Exception:
                    # Vẫn không đọc được rich text -> mở chế độ thường
                    self.rich_text_ok = False
                    self.wb = openpyxl.load_workbook(repaired, data_only=False)
            except Exception as e2:
                self._cleanup_tmp()
                raise ValueError(
                    "Không đọc được nội dung file Excel dù đã thử sửa chữa.\n"
                    "File có thể bị hỏng nặng hoặc định dạng không tương thích."
                ) from e2
        if CHECKLIST not in self.wb.sheetnames:
            raise ValueError(
                f"File không có sheet '{CHECKLIST}'.\n"
                "Hãy chọn đúng file checklist 52 tiêu chí của HR audit (kèm sheet Checklist)."
            )
        # Ảnh chụp nội dung cột L của FILE ĐẦU VÀO: dùng cho hậu kiểm "cột L đầy
        # đủ" (nội dung đánh giá của kiểm toán viên không được rơi mất khi app
        # trình bày lại khối song ngữ).
        self.original_evidence = {
            row: self.evidence(row) for row in sorted(ROW_ITEMS)
        }

    def _cleanup_tmp(self):
        """Xóa file tạm (bản sao đã sửa chữa) nếu có."""
        if self._tmp:
            try:
                os.remove(self._tmp)
            except OSError:
                pass
            self._tmp = None

    # ------------------------------------------------------------------
    # Checklist
    # ------------------------------------------------------------------
    def audit_id(self, row):
        return ROW_ITEMS.get(row)

    def rating(self, row):
        return self.wb[CHECKLIST].cell(row=row, column=4).value

    def evidence(self, row):
        """
        Nội dung cột L dưới dạng chuỗi thuần (CellRichText được flatten
        để parse/GUI xử lý như text bình thường).

        Chuẩn hóa Unicode dạng NFC để mọi phép so khớp header khối (in đậm, tách
        ngôn ngữ, chèn dòng phân cách) luôn đúng bất kể file đầu vào lưu tiếng
        Việt ở dạng NFD (decomposed) hay NFC (composed). NFC và NFD hiển thị
        GIỐNG NHAU trong Excel nên việc chuẩn hóa này không làm thay đổi nội
        dung hiển thị, chỉ đảm bảo nhận diện header "NỘI DUNG ĐÁNH GIÁ
        (TIẾNG VIỆT)" để in đậm.
        """
        v = self.wb[CHECKLIST].cell(row=row, column=12).value
        text = str(v) if isinstance(v, CellRichText) else v
        if isinstance(text, str):
            text = _norm_unicode(text)
        return text

    def sub_category(self, row):
        return self.wb[CHECKLIST].cell(row=row, column=10).value

    def set_rating(self, row, rating):
        if rating not in ALLOWED_RATINGS:
            raise ValueError(f"Rating không hợp lệ: {rating}")
        self.wb[CHECKLIST].cell(row=row, column=4).value = rating

    def set_evidence(self, row, text):
        cell = self.wb[CHECKLIST].cell(row=row, column=12)
        if self.rich_text_ok:
            # Header khối song ngữ được in đậm bằng rich text (đúng file mẫu).
            # Mọi run mang rPr đầy đủ (kế thừa font của ô) để Excel hiển thị chuẩn.
            # Luôn đảm bảo đúng 1 hàng trống ngăn cách phần EN và VI bất kể
            # nội dung đến từ đâu (build_content đã làm, nhưng set_evidence là
            # chốt chặn cuối cùng trước khi ghi vào ô).
            if isinstance(text, str) and HEADER_VI in text:
                text = ensure_section_gap(text)
            cell.value = to_rich_text(text, base_font=cell.font)
        else:
            cell.value = text
        self._auto_fit_row(row, text)

    def _auto_fit_row(self, row, text):
        """
        Tự động tăng chiều cao dòng và căn lề trên cho ô L khi nội dung dài
        (mô phỏng thao tác trình bày của người biên tập trong file mẫu).
        """
        if text is not None and not isinstance(text, str):
            text = str(text)
        if not isinstance(text, str) or len(text) < 300:
            return
        ws = self.wb[CHECKLIST]
        cell = ws.cell(row=row, column=12)
        align = copy.copy(cell.alignment)
        align.vertical = "top"
        align.wrap_text = True
        cell.alignment = align

        width = ws.column_dimensions["L"].width or 40.0
        cpl = max(20, int(width))  # số ký tự mỗi dòng ước tính
        est_lines = 0
        for ln in text.split("\n"):
            est_lines += max(1, -(-len(ln) // cpl))
        new_height = est_lines * 13.5 + 6
        rd = ws.row_dimensions.get(row)
        old_height = rd.height if rd is not None and rd.height else 0
        if new_height > old_height:
            ws.row_dimensions[row].height = new_height


    def rows(self):
        """Danh sách dict (row, audit_id, rating, sub_category) cho 52 dòng tiêu chí."""
        return [
            {
                "row": row,
                "audit_id": ROW_ITEMS[row],
                "rating": self.rating(row),
                "sub_category": self.sub_category(row),
            }
            for row in sorted(ROW_ITEMS)
        ]

    def normalize_empty_rows(self):
        """
        Chuẩn hóa các dòng N/A đang trống hoặc placeholder cũ
        về placeholder chuẩn của báo cáo (giống file đầu ra được duyệt).

        Yêu cầu: dòng có Final Significance (cột D) = Conformance thì KHÔNG
        được đụng tới cột L - giữ nguyên văn từ file đầu vào (kể cả khi trống).
        """
        changed = []
        for row in sorted(ROW_ITEMS):
            rating = (self.rating(row) or "").strip()
            if rating == "N/A":
                if is_placeholder(self.evidence(row)):
                    self.set_evidence(row, PLACEHOLDER)
                    changed.append(row)
        return changed

    def repair_bilingual_formatting(self):
        """
        Rà soát lại cột L của các ô đã có khối song ngữ:
        - Nếu nội dung lẫn ngôn ngữ (EN còn tiếng Việt / VI còn tiếng Anh) -> tách
          thành EN thuần Anh / VI thuần Việt (dịch bù cụm máy dịch bỏ sót);
        - Đảm bảo đúng 1 hàng trống ngăn cách phần EN và VI;
        - Đảm bảo 4 header khối (AUDIT FINDINGS (ENGLISH) / Legal references: /
          NỘI DUNG ĐÁNH GIÁ (TIẾNG VIỆT) / Căn cứ pháp lý:) đều được in đậm.

        Áp dụng cho MỌI dòng (kể cả Conformance): dòng Conformance chỉ được sửa
        ĐỊNH DẠNG/NGÔN NGỮ để trình bày đúng mẫu - không thêm trích dẫn KB và
        không thay nội dung đánh giá (nhờ add_kb_refs=False).

        Chỉ rebuild lại khối khi THỰC SỰ có mục lẫn ngôn ngữ; các dòng đã thuần
        (vd nội dung curated đã duyệt) chỉ được sửa hàng trống + in đậm, văn bản
        giữ nguyên văn. Dòng placeholder/draft thường không bị đụng tới.
        Mục đích: mở file cũ (phiên bản app trước) rồi lưu lại (GUI/CLI) thì nội
        dung ra đúng định dạng mẫu mà không làm xê dịch ô đã chuẩn.
        Trả về danh sách dòng đã sửa.
        """
        ws = self.wb[CHECKLIST]
        changed = []
        for row in sorted(ROW_ITEMS):
            cell = ws.cell(row=row, column=12)
            v = cell.value
            text = self.evidence(row)
            if not isinstance(text, str) or not text.strip():
                continue
            headers = [h for h in BILINGUAL_HEADERS if h in text]
            if not headers:
                continue  # placeholder / draft thường - không đụng tới

            old_text = text

            # Nội dung THỰC SỰ lẫn ngôn ngữ -> rebuild khối chuẩn; khối đã thuần
            # thì giữ nguyên văn (chỉ rà hàng trống + in đậm bên dưới).
            if HEADER_EN in text and HEADER_VI in text:
                parts = parse_bilingual(text)
                en_f, en_r = parts["en_findings"], parts["en_refs"]
                vi_f, vi_r = parts["vi_findings"], parts["vi_refs"]
                pur_en, pur_vi = purify_findings(list(en_f), list(vi_f))
                lang_dirty = (pur_en != list(en_f)) or (pur_vi != list(vi_f))
                # phần trích dẫn ghi lẫn ngôn ngữ cũng tính là "dirty"
                if not lang_dirty:
                    _, sep_en_r, _, sep_vi_r = separate_languages(
                        [], list(en_r), [], list(vi_r))
                    lang_dirty = (sep_en_r != list(en_r)) or (sep_vi_r != list(vi_r))
                if not lang_dirty:
                    lang_dirty = bool(scan_language_mixing(text))
                if lang_dirty:
                    text = normalize_bilingual_block(
                        ROW_ITEMS[row], text, add_kb_refs=False)

            # 1) Hàng trống ngăn cách EN/VI
            text = ensure_section_gap(text)

            # 2) 4 header khối có được in đậm đầy đủ không?
            bold_ok = False
            if isinstance(v, CellRichText):
                bold_lines = {
                    _norm_unicode(blk.text).rstrip("\n").strip()
                    for blk in v
                    if hasattr(blk, "text") and hasattr(blk, "font")
                    and getattr(blk.font, "b", False)
                }
                bold_ok = all(h in bold_lines for h in headers)

            if text == old_text and ensure_section_gap(old_text) == old_text and bold_ok:
                continue
            self.set_evidence(row, text)
            changed.append(row)
        return changed

    def regenerate_findings(self, curated=None, use_curated=True, add_kb_refs=True,
                            apply_ratings=None):
        """
        Sinh lại nội dung cột L cho TẤT CẢ 52 dòng theo ĐÚNG MẪU BÁO CÁO
        (khối song ngữ 4 header in đậm: AUDIT FINDINGS (ENGLISH) + Legal references:
        / NỘI DUNG ĐÁNH GIÁ (TIẾNG VIỆT) + Căn cứ pháp lý:), với phần tiếng Anh
        THUẦN ANH và phần tiếng Việt THUẦN VIỆT (không còn lẫn Anh/Việt).

        Quy tắc nội dung:
        - Dòng vi phạm (Final Significance khác Conformance): ưu tiên nội dung
          curated đã kiểm chuẩn; nếu không có thì giữ draft và bổ sung trích dẫn KB.
        - Dòng Conformance: GIỮ NGUYÊN nội dung draft của file đầu vào (không ghi
          đè bằng curated, không bổ sung trích dẫn KB cho phần có draft) và KHÔNG
          đổi đánh giá ở cột D - chỉ được TRÌNH BÀY LẠI theo khối mẫu (tách ngôn
          ngữ, không để lẫn Anh/Việt).
        - Draft trống: khối rỗng + trích dẫn KB (cả dòng Conformance).

        apply_ratings: dict {row: rating} - ghi đè rating của curated nếu có.
        Nếu curated entry chứa "rating" thì cập nhật luôn cột D (đánh giá mức độ
        do người duyệt quyết định, lưu cùng nội dung curated) - CHỈ áp dụng cho
        dòng đang không phải Conformance; dòng Conformance trong file đầu vào
        giữ nguyên đánh giá (trừ khi người dùng chỉ định tường minh qua
        apply_ratings). Trả về danh sách dòng đã xử lý.
        """
        if apply_ratings:
            for row, rating in apply_ratings.items():
                self.set_rating(row, rating)
        touched = []
        for row in sorted(ROW_ITEMS):
            rating = (self.rating(row) or "").strip()
            audit_id = ROW_ITEMS[row]
            cur = (curated or {}).get(audit_id) if use_curated else None
            # Rating curated chỉ áp dụng cho dòng đang KHÔNG phải Conformance;
            # dòng Conformance trong file đầu vào giữ nguyên đánh giá (trừ khi
            # người dùng chỉ định tường minh qua apply_ratings).
            if (cur and cur.get("rating") and rating != "Conformance"
                    and not (apply_ratings and row in apply_ratings)):
                self.set_rating(row, cur["rating"])
                rating = cur["rating"]
            is_conformance = rating == "Conformance"
            draft = self.evidence(row)
            # Dòng Conformance: không dùng curated và không bổ sung trích dẫn KB
            # cho phần có draft -> nội dung đánh giá giữ nguyên từ file đầu vào.
            content = build_content(
                audit_id, draft, rating=rating,
                curated=({} if is_conformance else curated),
                use_curated=(use_curated and not is_conformance),
                add_kb_refs=(add_kb_refs and not is_conformance),
            )
            self.set_evidence(row, content)
            touched.append(row)
        return touched

    def language_mixing_report(self):
        """
        Hậu kiểm toàn bộ cột L: liệt kê mọi mục còn LẪN NGÔN NGỮ (EN còn tiếng
        Việt, VI còn tiếng Anh) hoặc ô KHÔNG theo khối mẫu.

        Trả về [(row, kind, content), ...] với kind = 'en' | 'vi' | 'format'.
        Danh sách rỗng = cột L đã đúng mẫu và sạch ngôn ngữ.
        """
        found = []
        for row in sorted(ROW_ITEMS):
            for kind, content in scan_language_mixing(self.evidence(row)):
                found.append((row, kind, content))
        return found

    def completeness_report(self):
        """
        Hậu kiểm "cột L ĐẦY ĐỦ": liệt kê mọi nội dung của FILE ĐẦU VÀO (cột L) không
        còn xuất hiện trong nội dung hiện tại của ô (sau khi Auto-generate/trình bày
        lại khối song ngữ).

        Trả về [(row, missing_content), ...]. Danh sách rỗng = cột L vẫn giữ đủ nội
        dung đánh giá của kiểm toán viên.
        """
        found = []
        for row in sorted(ROW_ITEMS):
            for item in missing_input_items(self.evidence(row),
                                            self.original_evidence.get(row)):
                found.append((row, item))
        return found

    # ------------------------------------------------------------------
    # Metadata đợt audit
    # ------------------------------------------------------------------
    def set_checklist_date(self, text):
        """Ô Checklist!D1 - ngày audit hiển thị trên báo cáo."""
        self.wb[CHECKLIST].cell(row=1, column=4).value = text

    def set_general_info(self, audit_type=None, team_leader=None, team_member=None,
                         closing_date=None, worker_interview_count=None,
                         initial_start=None, initial_end=None):
        gi = self.wb[GENERAL_INFO]
        if team_leader is not None:
            gi.cell(row=6, column=4).value = team_leader
        if team_member is not None:
            gi.cell(row=7, column=4).value = team_member
        if audit_type is not None:
            gi.cell(row=8, column=4).value = audit_type
        if initial_start is not None:
            gi.cell(row=9, column=4).value = initial_start
        if initial_end is not None:
            gi.cell(row=10, column=4).value = initial_end
        if closing_date is not None:
            gi.cell(row=11, column=4).value = closing_date
        if worker_interview_count is not None:
            gi.cell(row=31, column=4).value = worker_interview_count

    def set_company_profile(self, year=None, interview_score=None, clear_old=True):
        cp = self.wb[COMPANY_PROFILE]
        if year is not None:
            cp.cell(row=2, column=6).value = f"Percentage of dealing with SEV/T in {year}"
            cp.cell(row=3, column=6).value = f"Total Sales amount in {year}"
        if interview_score is not None:
            cp.cell(row=11, column=6).value = interview_score
        if clear_old:
            for coord in PROFILE_RESET_CELLS:
                col = openpyxl.utils.column_index_from_string(coord[0])
                row = int(coord[1:])
                cell = cp.cell(row=row, column=col)
                cell.value = None
                cell.hyperlink = None  # xóa cả hyperlink cũ (vd mailto) nếu còn sót

    # ------------------------------------------------------------------
    # Lưu
    # ------------------------------------------------------------------
    def save(self, path=None):
        target = path or self.path
        # Đánh dấu phiên bản app vào thuộc tính file để kiểm tra provenance
        # (Excel: File > Info > Properties > Tags = "HR-Audit-App v3.2")
        try:
            self.wb.properties.keywords = f"HR-Audit-App v{APP_VERSION}"
        except Exception:
            pass
        self.wb.save(target)
        return target

    def close(self):
        try:
            self.wb.close()
        except Exception:
            pass
        self._cleanup_tmp()

